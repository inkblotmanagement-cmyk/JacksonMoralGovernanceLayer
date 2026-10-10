"""Shared helpers for reading the two labelers' filled sheets."""
from __future__ import annotations
import csv
from pathlib import Path

KIT = Path(__file__).resolve().parent
LABELS = ["Safe - just help", "Fix or limit it first", "Refuse - it's harmful",
          "Get a human / crisis help"]
LABEL_TO_DECISION = {LABELS[0]: "ALLOW", LABELS[1]: "MODIFY",
                     LABELS[2]: "BLOCK", LABELS[3]: "ESCALATE"}


def _norm_label(v):
    if v is None:
        return None
    s = " ".join(str(v).split()).strip()
    if not s:
        return None
    if s in LABEL_TO_DECISION:
        return s
    # tolerate minor variants / people typing the decision directly
    low = s.lower()
    for lab, dec in LABEL_TO_DECISION.items():
        if low == lab.lower() or low == dec.lower():
            return lab
    for lab in LABELS:
        if low.startswith(lab.split(" - ")[0].lower().split(" ")[0]):
            pass
    aliases = {"safe": LABELS[0], "allow": LABELS[0], "fix": LABELS[1], "limit": LABELS[1],
               "modify": LABELS[1], "refuse": LABELS[2], "block": LABELS[2], "harmful": LABELS[2],
               "human": LABELS[3], "crisis": LABELS[3], "escalate": LABELS[3], "unsure": LABELS[3]}
    for k, lab in aliases.items():
        if k in low:
            return lab
    return s  # unknown -> returned as-is so the caller can flag it


def read_filled(path):
    """Read one labeler's filled file (.xlsx or .csv, single or multi-session).
    Returns {item_id: {"label": <canonical label or raw>, "notes": str}}."""
    path = Path(path)
    out = {}
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True)
        for ws in wb.worksheets:
            rows = list(ws.iter_rows(values_only=True))
            if not rows or not rows[0]:
                continue
            hdr = [str(c).strip().lower() if c else "" for c in rows[0]]
            if not any(h.startswith("item_id") for h in hdr):
                continue
            idx = {h: i for i, h in enumerate(hdr)}
            ii = next(i for h, i in idx.items() if h.startswith("item_id"))
            li = next((i for h, i in idx.items() if h.startswith("label")), None)
            ni = next((i for h, i in idx.items() if h.startswith("notes")), None)
            for r in rows[1:]:
                if not r or ii >= len(r) or not r[ii]:
                    continue
                iid = str(r[ii]).strip()
                lab = _norm_label(r[li]) if li is not None and li < len(r) else None
                notes = (str(r[ni]).strip() if ni is not None and ni < len(r) and r[ni] else "")
                out[iid] = {"label": lab, "notes": notes}
    else:
        with path.open() as f:
            rd = csv.reader(f)
            rows = list(rd)
        hdr = [h.strip().lower() for h in rows[0]]
        ii = next(i for i, h in enumerate(hdr) if h.startswith("item_id"))
        li = next((i for i, h in enumerate(hdr) if h.startswith("label")), None)
        ni = next((i for i, h in enumerate(hdr) if h.startswith("notes")), None)
        for r in rows[1:]:
            if not r or not r[ii]:
                continue
            out[r[ii].strip()] = {"label": _norm_label(r[li]) if li is not None else None,
                                  "notes": (r[ni].strip() if ni is not None and ni < len(r) else "")}
    return out


def read_answer_key():
    out = {}
    with (KIT / "_answer_key.csv").open() as f:
        for row in csv.DictReader(f):
            out[row["item_id"]] = row
    return out
